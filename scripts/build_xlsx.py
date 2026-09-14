#!/usr/bin/env python3
"""Local Excel export of the library: dist/Design_Arsenal.xlsx.

  python3 scripts/build_xlsx.py [--out PATH]

Every text cell is written as a literal string (never a formula), so third-party text can't
execute in a spreadsheet app. Summary counts on "Start Here" are live formulas.
"""
from __future__ import annotations

import argparse
import sys
from collections import OrderedDict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib import store  # noqa: E402

from openpyxl import Workbook  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.workbook.properties import CalcProperties  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F2430")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BAND = Font(name=FONT, bold=True, color="1F2430", size=9)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
LINK = Font(name=FONT, size=10, color="1155CC", underline="single")
TITLE = Font(name=FONT, size=18, bold=True, color="1F2430")
H2 = Font(name=FONT, size=12, bold=True, color="1F2430")
MUTED = Font(name=FONT, size=9, italic=True, color="6B7280")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(bottom=Side(style="thin", color="E5E7EB"))
TIER_COLORS = {"Core": "C6EFCE", "Strong": "DDEBF7", "Situational": "FFF2CC", "Low fit": "EDEDED"}
PRICE_COLORS = {"Free": "C6EFCE", "Open source (free)": "C6EFCE", "Freemium": "E2F0D9",
                "Free trial then paid": "FCE4D6", "Paid": "F8CBAD", "Unconfirmed": "EDEDED"}
TIER_ORDER = {"Core": 0, "Strong": 1, "Situational": 2, "Low fit": 3}


def text(cell, value):
    """Write a value; strings are forced to literal text."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def j(v):
    return ", ".join(v) if isinstance(v, list) else ("" if v is None else v)


def glance(r):
    m, mo, ot = r["pricing_model"], r["paid_entry_usd_month"], r["one_time_usd"]
    if m in ("Free", "Open source (free)"):
        return "Free"
    parts = [f"${mo:g}/mo"] if mo else []
    if ot:
        parts.append(f"${ot:g} one-time")
    if not parts:
        parts = ["price unconfirmed"] if m != "Paid" else ["quote-based / unconfirmed"]
    lead = "Free tier + " if m == "Freemium" else ("Trial, then " if m.startswith("Free trial") else "")
    return lead + " or ".join(parts)


def header(ws, row, headers, widths=None):
    for c, h in enumerate(headers, 1):
        cell = text(ws.cell(row=row, column=c), h)
        cell.font, cell.fill = HDR_FONT, HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 32
    for c, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(c)].width = w


COLS = [
    ("IDENTITY", "ID", 16, lambda r: r["id"]),
    ("IDENTITY", "Name", 20, lambda r: r["name"]),
    ("IDENTITY", "URL", 28, lambda r: r["url"]),
    ("PORTFOLIO FIT", "Category", 22, lambda r: r["category"]),
    ("PORTFOLIO FIT", "Tier", 11, lambda r: r["tier"]),
    ("PORTFOLIO FIT", "Portfolio fit (1-5)", 9, lambda r: r["portfolio_fit"]),
    ("PORTFOLIO FIT", "One-liner", 34, lambda r: r["one_liner"]),
    ("PORTFOLIO FIT", "Best use for a portfolio", 46, lambda r: r["best_use_for_portfolio"]),
    ("PORTFOLIO FIT", "Portfolio jobs", 26, lambda r: j(r["portfolio_jobs"])),
    ("PORTFOLIO FIT", "Client work OK?", 14, lambda r: r["client_ok"]),
    ("PORTFOLIO FIT", "Why this tier", 36, lambda r: r["tier_reason"]),
    ("WHAT IT IS", "What it offers", 50, lambda r: r["what_it_offers"]),
    ("WHAT IT IS", "Subcategory", 20, lambda r: r["subcategory"]),
    ("WHAT IT IS", "Output", 14, lambda r: r["output_type"]),
    ("WHAT IT IS", "Tech stack", 22, lambda r: r["tech_stack"]),
    ("WHAT IT IS", "Skill level", 11, lambda r: r["skill_level"]),
    ("COST", "Price at a glance", 22, glance),
    ("COST", "Pricing model", 14, lambda r: r["pricing_model"]),
    ("COST", "Free tier", 34, lambda r: r["free_tier"]),
    ("COST", "Cheapest paid ($/mo)", 11, lambda r: r["paid_entry_usd_month"]),
    ("COST", "Cheapest paid (detail)", 30, lambda r: r["paid_entry_label"]),
    ("COST", "One-time ($)", 10, lambda r: r["one_time_usd"]),
    ("COST", "Commercial use / license", 34, lambda r: r["commercial_use_license"]),
    ("COST", "Client-use note", 30, lambda r: r["client_ok_note"]),
    ("PRACTICAL", "AI-agent access", 30, lambda r: r["agent_friendly"]),
    ("PRACTICAL", "Maintenance signal", 28, lambda r: r["maintenance_signal"]),
    ("PRACTICAL", "Overlap group", 18, lambda r: r["overlap_group"]),
    ("PRACTICAL", "Pick it over overlaps when…", 34, lambda r: r["pick_over_overlaps"]),
    ("PRACTICAL", "Caveats", 34, lambda r: r["caveats"]),
    ("TRUST", "Status", 9, lambda r: r["status"]),
    ("TRUST", "Last verified", 11, lambda r: r["last_verified_at"]),
    ("TRUST", "Confidence", 10, lambda r: r["confidence"]),
    ("TRUST", "Unverified fields", 22, lambda r: j(r["unverified_fields"])),
    ("TRUST", "Pricing source", 30, lambda r: r["pricing_source_url"]),
    ("TRUST", "Notes", 30, lambda r: r["user_notes"]),
]


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(store.ROOT / "dist" / "Design_Arsenal.xlsx"))
    out = Path(ap.parse_args().out)

    resources = list(store.load_state("resources").values())
    inbox = list(store.load_state("inbox").values())
    tax = store.taxonomy()
    gaps = store.read_json(store.CONFIG / "gaps.json")["gaps"]
    leads = store.read_jsonl(store.DATA / "leads.jsonl")
    checks = store.read_jsonl(store.DATA / "provenance" / "seed_fact_checks.jsonl")
    names = {r["id"]: r["name"] for r in resources + inbox}

    wb = Workbook()
    ws = wb.active
    ws.title = "Resources"
    head = [c[1] for c in COLS]
    col = {h: i + 1 for i, h in enumerate(head)}
    prev = None
    for i, (band, *_r) in enumerate(COLS, 1):
        if band != prev:
            text(ws.cell(row=1, column=i), band).font = BAND
            prev = band
    header(ws, 2, head, [c[2] for c in COLS])
    rows = sorted(resources, key=lambda r: (TIER_ORDER[r["tier"]], -r["portfolio_fit"], r["category"], r["name"].lower()))
    first = 3
    for n, r in enumerate(rows):
        row = first + n
        for c, (_b, _h, _w, get) in enumerate(COLS, 1):
            cell = text(ws.cell(row=row, column=c), get(r))
            cell.font, cell.alignment, cell.border = BODY, WRAP, THIN
        est = max(len(str(ws.cell(row=row, column=col[h]).value or "")) / COLS[col[h] - 1][2]
                  for h in ("One-liner", "Best use for a portfolio", "Why this tier"))
        ws.row_dimensions[row].height = min(max(15 * (est + 1), 45), 120)
        u = ws.cell(row=row, column=col["URL"])
        u.hyperlink, u.font = r["url"], LINK
        ws.cell(row=row, column=col["Name"]).font = BOLD
        ws.cell(row=row, column=col["Cheapest paid ($/mo)"]).number_format = "$#,##0.00;;-"
        ws.cell(row=row, column=col["One-time ($)"]).number_format = "$#,##0;;-"
    last = first + max(len(rows), 1) - 1
    ws.freeze_panes = ws.cell(row=first, column=col["Category"])
    ws.auto_filter.ref = f"A2:{get_column_letter(len(COLS))}{last}"
    for h, palette in (("Tier", TIER_COLORS), ("Pricing model", PRICE_COLORS)):
        L = get_column_letter(col[h])
        for val, color in palette.items():
            ws.conditional_formatting.add(f"{L}{first}:{L}{last}", FormulaRule(
                formula=[f'${L}{first}="{val}"'], fill=PatternFill("solid", fgColor=color)))
    dv = DataValidation(type="list", formula1='"Core,Strong,Situational,Low fit"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{get_column_letter(col['Tier'])}{first}:{get_column_letter(col['Tier'])}{last + 200}")
    rng = lambda h: f"Resources!${get_column_letter(col[h])}${first}:${get_column_letter(col[h])}${last}"  # noqa: E731

    ov = wb.create_sheet("Start Here", 0)
    for L, w in zip("ABCDEF", (34, 10, 10, 12, 14, 64)):
        ov.column_dimensions[L].width = w
    text(ov["A1"], "Design Arsenal").font = TITLE
    text(ov["A2"], f"{len(resources)} resources for building portfolio websites · {len(inbox)} awaiting review · "
                   f"{sum(1 for x in leads if x['status'] == 'queued')} queued leads").font = MUTED
    r = 4
    text(ov.cell(row=r, column=1), "By category").font = H2
    r += 1
    header(ov, r, ["Category", "Sites", "Core", "Avg fit", "Client OK", "Definition"])
    start = r + 1
    for t in tax["categories"]:
        r += 1
        text(ov.cell(row=r, column=1), t["category"]).font = BOLD
        A = f"$A{r}"
        ov.cell(row=r, column=2, value=f"=COUNTIF({rng('Category')},{A})")
        ov.cell(row=r, column=3, value=f'=COUNTIFS({rng("Category")},{A},{rng("Tier")},"Core")')
        ov.cell(row=r, column=4, value=f'=IFERROR(AVERAGEIF({rng("Category")},{A},{rng("Portfolio fit (1-5)")}),"-")').number_format = "0.0"
        ov.cell(row=r, column=5, value=f'=COUNTIFS({rng("Category")},{A},{rng("Client work OK?")},"Yes")')
        text(ov.cell(row=r, column=6), t["definition"]).alignment = WRAP
        ov.row_dimensions[r].height = 30
    r += 1
    text(ov.cell(row=r, column=1), "Total").font = BOLD
    for c, L in ((2, "B"), (3, "C"), (5, "E")):
        ov.cell(row=r, column=c, value=f"=SUM({L}{start}:{L}{r - 1})").font = BOLD
    r += 2
    text(ov.cell(row=r, column=1), "By tier").font = H2
    for t, why in tax["tiers"].items():
        r += 1
        text(ov.cell(row=r, column=1), t).font = BOLD
        ov.cell(row=r, column=2, value=f"=COUNTIF({rng('Tier')},$A{r})")
        text(ov.cell(row=r, column=6), why)
    r += 2
    text(ov.cell(row=r, column=1), "Cost picture").font = H2
    for pm in ["Free", "Open source (free)", "Freemium", "Free trial then paid", "Paid", "Unconfirmed"]:
        r += 1
        text(ov.cell(row=r, column=1), pm).font = BOLD
        ov.cell(row=r, column=2, value=f"=COUNTIF({rng('Pricing model')},$A{r})")
    r += 1
    text(ov.cell(row=r, column=1), "Sum of cheapest paid plans ($/mo)").font = BOLD
    ov.cell(row=r, column=2, value=f"=SUM({rng('Cheapest paid ($/mo)')})").number_format = "$#,##0"
    r += 1
    text(ov.cell(row=r, column=1), "…Core tier only ($/mo)").font = BOLD
    ov.cell(row=r, column=2, value=f'=SUMIF({rng("Tier")},"Core",{rng("Cheapest paid ($/mo)")})').number_format = "$#,##0"

    jobs = wb.create_sheet("By Portfolio Job")
    header(jobs, 1, ["Portfolio job", "Resource", "Tier", "Fit", "Price", "Best use"], [30, 24, 12, 6, 22, 80])
    i = 2
    for job in tax["portfolio_jobs"]:
        members = sorted((x for x in resources if job in x["portfolio_jobs"]),
                         key=lambda x: (TIER_ORDER[x["tier"]], -x["portfolio_fit"]))
        for m in members:
            for c, v in enumerate([job, m["name"], m["tier"], m["portfolio_fit"], glance(m), m["best_use_for_portfolio"]], 1):
                cell = text(jobs.cell(row=i, column=c), v)
                cell.font, cell.alignment = BODY, WRAP
            text(jobs.cell(row=i, column=1), job).font = BOLD
            i += 1
        i += 1 if members else 0
    jobs.freeze_panes = "A2"

    ol = wb.create_sheet("Overlaps")
    header(ol, 1, ["Overlap group", "Resource", "Tier", "Price", "Pick it when…"], [28, 22, 12, 22, 80])
    groups = OrderedDict()
    for x in sorted(resources, key=lambda x: x["overlap_group"]):
        if x["overlap_group"] and x["overlap_group"].lower() != "unique":
            groups.setdefault(x["overlap_group"], []).append(x)
    i = 2
    for g, members in groups.items():
        if len(members) < 2:
            continue
        for m in sorted(members, key=lambda m: TIER_ORDER[m["tier"]]):
            for c, v in enumerate([g, m["name"], m["tier"], glance(m), m["pick_over_overlaps"]], 1):
                cell = text(ol.cell(row=i, column=c), v)
                cell.font, cell.alignment = BODY, WRAP
            i += 1
        i += 1
    ol.freeze_panes = "A2"

    gp = wb.create_sheet("Gaps")
    header(gp, 1, ["What the library is missing", "Why it matters", "Search queries"], [34, 60, 70])
    for i, g in enumerate(gaps, 2):
        for c, v in enumerate([g["gap"], g["why_it_matters"], "\n".join(g["queries"])], 1):
            cell = text(gp.cell(row=i, column=c), v)
            cell.font, cell.alignment = BODY, WRAP

    ld = wb.create_sheet("Leads")
    header(ld, 1, ["Lead", "URL", "Status", "Fit", "Kind", "One-liner", "Similar to", "Reason"], [22, 30, 10, 6, 20, 40, 30, 50])
    for i, L in enumerate(leads, 2):
        vals = [L["name"], L["url"], L["status"], L["portfolio_fit"], L.get("kind", ""), L.get("one_liner", ""),
                ", ".join(names.get(s, s) for s in L.get("similar_to", [])), L.get("reason", "")]
        for c, v in enumerate(vals, 1):
            cell = text(ld.cell(row=i, column=c), v)
            cell.font, cell.alignment = BODY, WRAP
        ld.cell(row=i, column=2).hyperlink = L["url"]
    ld.freeze_panes = "A2"
    if leads:
        ld.auto_filter.ref = f"A1:H{len(leads) + 1}"

    cr = wb.create_sheet("Fact-check Log")
    header(cr, 1, ["Resource", "Field", "First pass said", "Corrected to", "Evidence"], [22, 22, 40, 40, 50])
    for i, c_ in enumerate(checks, 2):
        for c, v in enumerate([names.get(c_["record_id"], c_["record_id"]), c_["field"], c_["first_pass"], c_["corrected"], c_["evidence"]], 1):
            cell = text(cr.cell(row=i, column=c), v)
            cell.font, cell.alignment = BODY, WRAP

    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out} resources={len(resources)} leads={len(leads)} checks={len(checks)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
